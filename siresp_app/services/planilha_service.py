"""
Serviço de leitura de planilhas de agenda (SIRESP).
Aceita .xls, .xlsx, .xlsm e .xls disfarçado de HTML.
Retorna lista de (profissional, especialidade, minutos).
"""
import os
import re
import unicodedata


def _normalizar(texto):
    """Remove acentos, colapsa espaços, upper e strip."""
    if texto is None:
        return ""
    texto = str(texto)
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = re.sub(r"\s+", " ", texto)
    return texto.strip().upper()


def ler_planilha(caminho):
    """
    Lê a planilha e retorna lista de (profissional, especialidade, minutos).
    Aceita .xlsx, .xlsm (openpyxl), .xls (xlrd) e .xls-HTML (pandas).
    """
    ext = os.path.splitext(caminho)[1].lower()

    if ext in (".xlsx", ".xlsm"):
        return _ler_xlsx(caminho)

    if ext == ".xls":
        try:
            return _ler_xls(caminho)
        except Exception as e:
            msg = str(e)
            if "BOF record" in msg or "Unsupported format" in msg or "not a zip" in msg:
                return _ler_xls_html(caminho)
            raise

    raise Exception(f"Extensão não suportada: {ext}")


def _ler_xlsx(caminho):
    from openpyxl import load_workbook
    wb = load_workbook(caminho, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    cabecalho = rows[0]
    return _extrair(cabecalho, rows[1:])


def _ler_xls(caminho):
    try:
        import xlrd
    except ImportError:
        raise ImportError("Biblioteca 'xlrd' não instalada. Rode: pip install xlrd")

    book = xlrd.open_workbook(caminho)
    sheet = book.sheet_by_index(0)
    if sheet.nrows == 0:
        return []
    cabecalho = [sheet.cell_value(0, c) for c in range(sheet.ncols)]
    linhas = []
    for r in range(1, sheet.nrows):
        linhas.append([sheet.cell_value(r, c) for c in range(sheet.ncols)])
    return _extrair(cabecalho, linhas)


def _ler_xls_html(caminho):
    try:
        import pandas as pd
    except ImportError:
        raise ImportError(
            "Para ler .xls disfarçado de HTML, instale: pip install pandas lxml html5lib"
        )

    dfs = pd.read_html(caminho, header=None)
    if not dfs:
        raise Exception("Nenhuma tabela HTML encontrada no arquivo.")

    df = max(dfs, key=lambda d: d.shape[0] * d.shape[1])
    df = df.astype(str)

    linhas = df.values.tolist()
    if not linhas:
        return []

    cabecalho = linhas[0]
    linhas_dados = linhas[1:]

    cabecalho = ["" if c.lower() == "nan" else c for c in cabecalho]
    linhas_dados = [
        ["" if c.lower() == "nan" else c for c in row]
        for row in linhas_dados
    ]

    return _extrair(cabecalho, linhas_dados)


def _extrair(cabecalho, linhas_dados):
    """Localiza colunas e extrai (profissional, especialidade, minutos)."""
    cab_norm = [_normalizar(c) for c in cabecalho]

    def achar_col(nome_alvo):
        alvo = _normalizar(nome_alvo)
        for i, c in enumerate(cab_norm):
            if c == alvo:
                return i
        for i, c in enumerate(cab_norm):
            if alvo in c:
                return i
        return None

    idx_prof = achar_col("PROFISSIONAL")
    idx_esp = achar_col("ESPECIALIDADE")
    idx_int = achar_col("INTERVALO ENTRE AS CONSULTAS")

    if idx_prof is None or idx_esp is None or idx_int is None:
        raise Exception(
            "Não achei as colunas obrigatórias. Esperado: "
            "PROFISSIONAL, ESPECIALIDADE, INTERVALO ENTRE AS CONSULTAS. "
            f"Encontrado: {cabecalho}"
        )

    resultado = []
    for linha in linhas_dados:
        if not linha:
            continue
        if max(idx_prof, idx_esp, idx_int) >= len(linha):
            continue

        prof = linha[idx_prof]
        esp = linha[idx_esp]
        intervalo = linha[idx_int]

        if prof is None or esp is None or intervalo is None:
            continue

        prof = str(prof).strip()
        esp = str(esp).strip()
        if not prof or not esp:
            continue

        try:
            if isinstance(intervalo, str):
                intervalo = intervalo.strip().replace(",", ".")
            if intervalo == "" or intervalo is None:
                continue
            minutos = int(float(intervalo))
        except (ValueError, TypeError):
            continue

        if minutos <= 0:
            continue

        resultado.append((prof, esp, minutos))

    return resultado