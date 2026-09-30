"""
Resumo por faturamento: uma linha por equipe (soma dos profissionais) e uma
por profissional avulso. Usado no fim das telas, dos Excel e dos PDF.
"""


def linhas_resumo(grupos, nome_item, horas_item, valor_item, id_item):
    """
    grupos: saída de profissionais_service.agrupar_por_equipe.
    Retorna lista de dicts:
      {'nome', 'profissionais', 'horas', 'valor', 'equipe': bool,
       'equipe_id': int|None, 'item_id': int|None}
    """
    linhas = []
    for g in grupos:
        if g['equipe']:
            linhas.append({
                'nome': g['equipe'].nome,
                'profissionais': ', '.join(sorted({nome_item(i) for i in g['itens']})),
                'horas': g['horas'],
                'valor': g['valor'],
                'equipe': True,
                'equipe_id': g['equipe'].pk,
                'item_id': None,
            })
        else:
            for i in g['itens']:
                linhas.append({
                    'nome': nome_item(i),
                    'profissionais': nome_item(i),
                    'horas': horas_item(i),
                    'valor': valor_item(i),
                    'equipe': False,
                    'equipe_id': None,
                    'item_id': id_item(i),
                })
    return linhas


def anexar_resumo_planilha(ws, linhas, total_horas, total_valor,
                           col_nome=1, col_profs=2, col_horas=3, col_valor=4):
    """
    Escreve no fim da planilha o bloco "RESUMO POR FATURAMENTO": uma linha por
    equipe (nome da equipe, profissionais, horas e valor somado), uma por
    profissional avulso e o TOTAL GERAL DO REPASSE. Colunas 1-based.
    """
    from openpyxl.styles import Font, PatternFill, Alignment

    azul = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
    claro = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    verde = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    colunas = (col_nome, col_profs, col_horas, col_valor)

    linha = ws.max_row + 3          # deixa duas linhas em branco
    ws.cell(row=linha, column=col_nome, value="RESUMO POR FATURAMENTO (EQUIPES)").font = \
        Font(bold=True, size=13)

    linha += 1
    for col, txt in zip(colunas, ("Equipe / Profissional", "Profissionais", "Horas", "Valor (R$)")):
        c = ws.cell(row=linha, column=col, value=txt)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = azul
        c.alignment = Alignment(horizontal="center")

    for l in linhas:
        linha += 1
        ws.cell(row=linha, column=col_nome, value=l['nome'])
        ws.cell(row=linha, column=col_profs, value=l['profissionais'])
        ws.cell(row=linha, column=col_horas, value=float(l['horas'])).number_format = '0.00'
        ws.cell(row=linha, column=col_valor, value=float(l['valor'])).number_format = 'R$ #,##0.00'
        for col in colunas:
            c = ws.cell(row=linha, column=col)
            c.font = Font(bold=l['equipe'])
            if l['equipe']:
                c.fill = claro

    linha += 1
    ws.cell(row=linha, column=col_nome, value="TOTAL GERAL DO REPASSE")
    ws.cell(row=linha, column=col_horas, value=float(total_horas)).number_format = '0.00'
    ws.cell(row=linha, column=col_valor, value=float(total_valor)).number_format = 'R$ #,##0.00'
    for col in colunas:
        c = ws.cell(row=linha, column=col)
        c.font = Font(bold=True)
        c.fill = verde
