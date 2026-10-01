"""Unidade do SIRESP com a qual o login foi feito: status, gravação nas extrações e relatórios."""
import time
from io import BytesIO
from unittest import mock

from django.contrib.auth.models import User
from django.test import TransactionTestCase
from django.urls import reverse
from openpyxl import load_workbook

from .models import Extracao, ItemProducao, RegraMinuto, Repasse
from .services import repasse_service, scraper_service
from .services.repasse_service import competencia_padrao
from .services.resumo_service import texto_unidades
from .tests_papeis import PapeisBase
from .tests_unidade import ScraperFalso
from .tests_util import AguardaThreads

AME = {'valor': '2206_AME SAO JOSE DO RIO PRETO', 'nome': 'AME SAO JOSE DO RIO PRETO', 'codigo': '2206',
       'texto': '6056148 - AME SAO JOSE DO RIO PRETO - Administrador Unidade Reg'}
HOSP = {'valor': '2191_HOSP EST JOAO PAULO II', 'nome': 'HOSP EST JOAO PAULO II', 'codigo': '2191',
        'texto': '6236596 - HOSP EST JOAO PAULO II - Administrador Unidade Reg'}


class TextoUnidadesTests(PapeisBase):
    def test_texto_do_cabecalho(self):
        self.assertEqual(texto_unidades(['AME']), 'Unidade: AME')
        self.assertEqual(texto_unidades(['B', 'A', 'B', ' A ']), 'Unidades: A; B')
        self.assertEqual(texto_unidades(['', None, '  ']), 'Unidade: não informada')
        self.assertEqual(texto_unidades([]), 'Unidade: não informada')
        self.assertEqual(texto_unidades(['', 'AME']), 'Unidade: AME')      # vazios não contam


class UnidadeNoLoginTests(AguardaThreads, TransactionTestCase):
    def setUp(self):
        super().setUp()
        self.user = User.objects.create_user('medico', password='senha123')
        self.client.login(username='medico', password='senha123')
        scraper_service._SCRAPERS.clear()
        self.addCleanup(scraper_service._SCRAPERS.clear)

    def _logado(self, unidade):
        sc = ScraperFalso(aguardando=False)
        sc.logado = True
        sc.unidade_atual = unidade
        scraper_service._SCRAPERS[self.user.id] = sc
        return sc

    def test_status_logado_informa_a_unidade(self):
        self._logado(AME)
        r = self.client.get(reverse('siresp_app:producao_status_login')).json()
        self.assertEqual((r['estado'], r['logado']), ('logado', True))
        self.assertEqual(r['unidade']['nome'], 'AME SAO JOSE DO RIO PRETO')
        self.assertEqual(r['unidade']['codigo'], '2206')
        self.assertIn('Administrador Unidade Reg', r['unidade']['texto'])

    def test_logado_sem_pergunta_de_unidade_vem_nulo(self):
        self._logado(None)
        r = self.client.get(reverse('siresp_app:producao_status_login')).json()
        self.assertEqual(r['estado'], 'logado')
        self.assertIsNone(r['unidade'])

    def test_sem_sessao_nao_tem_unidade(self):
        r = self.client.get(reverse('siresp_app:producao_status_login')).json()
        self.assertEqual(r['estado'], 'sem_sessao')
        self.assertNotIn('unidade', r)
        self.assertIsNone(scraper_service.unidade_atual(self.user))

    def test_pagina_de_producao_tem_o_lugar_da_unidade(self):
        html = self.client.get(reverse('siresp_app:producao_home')).content.decode('utf-8')
        self.assertIn('id="status-unidade-nome"', html)

    def test_extracao_individual_grava_a_unidade_logada(self):
        self._logado(HOSP)
        dados = [{'Especialidade': 'CARDIOLOGIA', 'Oferta_N': '10'}]
        with mock.patch('siresp_app.services.scraper_service.extrair_producao', return_value=dados):
            r = self.client.post(reverse('siresp_app:producao_extrair'), data={
                'medico': {'nome': 'ANA LIMA', 'crm': '1', 'codigo': '9'},
                'data_ini': '01/08/2026', 'data_fim': '31/08/2026'},
                content_type='application/json').json()
            self.assertTrue(r['ok'], r)
            fim = time.time() + 5
            while time.time() < fim and not ItemProducao.objects.exists():
                time.sleep(0.05)
        ext = Extracao.objects.get()
        self.assertEqual((ext.unidade_nome, ext.unidade_codigo), ('HOSP EST JOAO PAULO II', '2191'))

    def test_extracao_em_lote_grava_a_unidade_logada(self):
        self._logado(AME)
        sc = 'siresp_app.services.lote_service.scraper_service'
        vp = 'siresp_app.views.producao.scraper_service'
        dados = [{'Especialidade': 'CARDIOLOGIA', 'Oferta_N': '10'}]
        with mock.patch(sc + '.listar_medicos', return_value=[{'nome': 'ANA LIMA', 'crm': '1', 'codigo': '1'}]), \
                mock.patch(sc + '.extrair_producao', return_value=dados), \
                mock.patch(vp + '.status_login', return_value={'estado': 'logado', 'logado': True}):
            r = self.client.post(reverse('siresp_app:producao_lote_iniciar'), data={
                'nomes': ['ANA LIMA'], 'data_ini': '01/08/2026', 'data_fim': '31/08/2026'},
                content_type='application/json').json()
            self.assertTrue(r['ok'], r)
            fim = time.time() + 8
            while time.time() < fim:
                st = self.client.get(reverse('siresp_app:producao_lote_status')).json()
                if not st['rodando']:
                    break
                time.sleep(0.1)
        ext = Extracao.objects.get()
        self.assertEqual((ext.unidade_nome, ext.unidade_codigo), ('AME SAO JOSE DO RIO PRETO', '2206'))

    def test_login_sem_unidade_grava_vazio_sem_quebrar(self):
        self._logado(None)
        dados = [{'Especialidade': 'CARDIOLOGIA', 'Oferta_N': '10'}]
        with mock.patch('siresp_app.services.scraper_service.extrair_producao', return_value=dados):
            self.client.post(reverse('siresp_app:producao_extrair'), data={
                'medico': {'nome': 'ANA LIMA'}, 'data_ini': '01/08/2026', 'data_fim': '31/08/2026'},
                content_type='application/json')
            fim = time.time() + 5
            while time.time() < fim and not ItemProducao.objects.exists():
                time.sleep(0.05)
        self.assertEqual(Extracao.objects.get().unidade_nome, '')


class UnidadeNosRelatoriosTests(PapeisBase):
    """Dois repasses na AME e um no hospital; um antigo sem unidade."""

    def setUp(self):
        super().setUp()
        RegraMinuto.objects.create(profissional_norm='ANA LIMA', especialidade_norm='CARDIOLOGIA', minutos=[30])
        self.exts = {}
        for nome, unid in (('ANA LIMA', AME), ('BRUNO DIAS', AME), ('CARLA MELO', HOSP), ('DIOGO SOLO', None)):
            RegraMinuto.objects.get_or_create(profissional_norm=nome, especialidade_norm='CARDIOLOGIA',
                                              defaults={'minutos': [30]})
            e = Extracao.objects.create(
                usuario_web=self.user, medico_nome=nome, data_ini='01/08/2026', data_fim='31/08/2026',
                unidade_nome=(unid or {}).get('nome', ''), unidade_codigo=(unid or {}).get('codigo', ''))
            ItemProducao.objects.create(extracao=e, especialidade='CARDIOLOGIA', ordem=0, dados={'Oferta_N': '40'})
            self.exts[nome] = e
        repasse_service.criar_repasses_lote(self.user, list(self.exts.values()))
        for r in Repasse.objects.all():
            repasse_service.finalizar_repasse(r, competencia_padrao())

    def ids(self, *nomes):
        return [Repasse.objects.get(extracao=self.exts[n]).pk for n in nomes]

    def _linhas(self, resp, aba=0):
        wb = load_workbook(BytesIO(resp.content))
        return [r for r in wb.worksheets[aba].iter_rows(values_only=True) if any(v is not None for v in r)]

    def test_excel_do_lote_em_nome_da_unidade(self):
        ids = ','.join(map(str, self.ids('ANA LIMA', 'BRUNO DIAS')))
        resp = self.client.get(reverse('siresp_app:repasse_lote_relatorio_excel') + '?ids=' + ids)
        linhas = self._linhas(resp)
        self.assertEqual(linhas[0][0], 'Relatório de Repasse Médico')
        self.assertTrue(linhas[1][0].startswith('Unidade: AME SAO JOSE DO RIO PRETO  |  Competência:'))
        cab = linhas[2]                                   # (linha em branco é descartada)
        self.assertEqual(cab[0], 'Equipe')
        self.assertEqual(cab[10], 'Unidade')
        dados = [r for r in linhas if r[3] == 'CARDIOLOGIA']
        self.assertEqual({r[10] for r in dados}, {'AME SAO JOSE DO RIO PRETO'})

    def test_excel_com_mais_de_uma_unidade_e_sem_unidade(self):
        ids = ','.join(map(str, self.ids('ANA LIMA', 'CARLA MELO')))
        r = self._linhas(self.client.get(reverse('siresp_app:repasse_lote_relatorio_excel') + '?ids=' + ids))
        self.assertTrue(r[1][0].startswith('Unidades: AME SAO JOSE DO RIO PRETO; HOSP EST JOAO PAULO II'))
        ids = ','.join(map(str, self.ids('DIOGO SOLO')))
        r = self._linhas(self.client.get(reverse('siresp_app:repasse_lote_relatorio_excel') + '?ids=' + ids))
        self.assertTrue(r[1][0].startswith('Unidade: não informada'))

    def test_excel_consolidado_e_individual(self):
        resp = self.client.post(reverse('siresp_app:relatorios_gerar_excel'),
                                {'repasses': self.ids('CARLA MELO')})
        wb = load_workbook(BytesIO(resp.content))
        ws = wb.worksheets[0]
        self.assertTrue(ws['A2'].value.startswith('Unidade: HOSP EST JOAO PAULO II'))
        self.assertEqual(ws.cell(row=4, column=18).value, 'Unidade')
        self.assertEqual(ws.cell(row=5, column=18).value, 'HOSP EST JOAO PAULO II')
        # individual
        pk = self.ids('ANA LIMA')[0]
        ind = load_workbook(BytesIO(self.client.get(reverse('siresp_app:repasse_excel', args=[pk])).content))
        self.assertIn('Unidade: AME SAO JOSE DO RIO PRETO', ind.worksheets[0]['A2'].value)

    def test_pdfs_geram_com_a_unidade(self):
        ids = ','.join(map(str, self.ids('ANA LIMA', 'CARLA MELO')))
        r = self.client.get(reverse('siresp_app:repasse_lote_relatorio_pdf') + '?ids=' + ids)
        self.assertEqual((r.status_code, r.content[:4]), (200, b'%PDF'))
        r = self.client.post(reverse('siresp_app:relatorios_gerar_pdf'), {'repasses': self.ids('ANA LIMA')})
        self.assertEqual((r.status_code, r.content[:4]), (200, b'%PDF'))
        # o subtítulo (com a unidade) chega ao gerador
        with mock.patch('siresp_app.services.pdf_service.gerar_pdf', return_value=b'%PDF-x') as g:
            self.client.post(reverse('siresp_app:relatorios_gerar_pdf'), {'repasses': self.ids('CARLA MELO')})
            self.assertIn('Unidade: HOSP EST JOAO PAULO II', g.call_args[0][1])
            self.client.get(reverse('siresp_app:repasse_lote_relatorio_pdf') + '?ids=' + ids)
            self.assertIn('Unidades: AME SAO JOSE DO RIO PRETO; HOSP EST JOAO PAULO II', g.call_args[0][1])

    def test_filtro_por_unidade_em_relatorios(self):
        url = reverse('siresp_app:relatorios_home')
        r = self.client.get(url)
        self.assertEqual(r.context['unidades'], ['AME SAO JOSE DO RIO PRETO', 'HOSP EST JOAO PAULO II'])
        self.assertEqual(len(r.context['repasses']), 4)
        self.assertContains(r, 'Todas as unidades')
        r = self.client.get(url + '?unidade=AME SAO JOSE DO RIO PRETO')
        self.assertEqual(sorted(x.extracao.medico_nome for x in r.context['repasses']), ['ANA LIMA', 'BRUNO DIAS'])
        r = self.client.get(url + '?unidade=OUTRA')
        self.assertEqual(r.context['repasses'], [])

    def test_telas_mostram_a_unidade_e_toleram_extracao_antiga(self):
        ana, diogo = (Repasse.objects.get(extracao=self.exts[n]) for n in ('ANA LIMA', 'DIOGO SOLO'))
        self.assertContains(self.client.get(reverse('siresp_app:repasse_ver', args=[ana.pk])), 'AME SAO JOSE DO RIO PRETO')
        self.assertContains(self.client.get(reverse('siresp_app:repasse_ver', args=[diogo.pk])), 'unidade não informada')
        self.assertContains(self.client.get(reverse('siresp_app:producao_ver_extracao', args=[self.exts['ANA LIMA'].pk])),
                            'AME SAO JOSE DO RIO PRETO')
        self.assertContains(self.client.get(reverse('siresp_app:producao_historico')), 'HOSP EST JOAO PAULO II')
        rel = self.client.get(reverse('siresp_app:repasse_lote_relatorio') + '?ids=%d' % ana.pk)
        self.assertContains(rel, 'Relatório em nome da unidade')
        self.assertContains(rel, 'Unidade: AME SAO JOSE DO RIO PRETO')
