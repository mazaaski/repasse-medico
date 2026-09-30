"""
Views do Repasse.
"""
import json
import math
from decimal import Decimal

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.views.decorators.http import require_POST, require_GET
from django.utils import timezone

from ..models import Extracao, Repasse, ItemRepasse
from ..services import repasse_service


# =========================================================
# ABRIR / CRIAR REPASSE
# =========================================================
@login_required
def abrir_repasse(request, extracao_pk):
    extracao = get_object_or_404(
        Extracao, pk=extracao_pk, usuario_web=request.user
    )

    repasse = Repasse.objects.filter(
        extracao=extracao, usuario_web=request.user
    ).first()

    if not repasse:
        repasse = repasse_service.criar_repasse_de_extracao(extracao, request.user)
        messages.success(request, 'Repasse criado a partir da extração.')

    return redirect('siresp_app:repasse_ver', pk=repasse.pk)


# =========================================================
# TELA DE REPASSE
# =========================================================
@login_required
def ver_repasse(request, pk):
    repasse = get_object_or_404(Repasse, pk=pk, usuario_web=request.user)
    itens = repasse.itens.all()

    ambiguos = _listar_ambiguidades(repasse)

    return render(request, 'siresp_app/repasse/home.html', {
        'repasse': repasse,
        'itens': itens,
        'qtd_ambiguos': len(ambiguos),
        'ambiguos': ambiguos,
    })


def _listar_ambiguidades(repasse):
    from ..models import RegraMinuto
    from ..services.repasse_service import _normalizar

    prof = repasse.extracao.medico_nome
    prof_n = _normalizar(prof)

    ambiguos = []
    for item in repasse.itens.all():
        esp_n = _normalizar(item.especialidade)
        try:
            r = RegraMinuto.objects.get(
                profissional_norm=prof_n,
                especialidade_norm=esp_n,
            )
            if len(r.minutos) > 1:
                ambiguos.append({
                    'item_id': item.id,
                    'esp': item.especialidade,
                    'valores': sorted(r.minutos),
                })
        except RegraMinuto.DoesNotExist:
            pass

    return ambiguos


# =========================================================
# TRAVA DE EDIÇÃO (helper)
# =========================================================
def _esta_finalizado(repasse):
    return repasse.status == 'finalizado'


# =========================================================
# ATUALIZAR ITEM
# =========================================================
@login_required
@require_POST
def atualizar_item(request, pk):
    item = get_object_or_404(
        ItemRepasse, pk=pk, repasse__usuario_web=request.user
    )

    if _esta_finalizado(item.repasse):
        return JsonResponse({
            'ok': False,
            'mensagem': 'Esse repasse está finalizado. Reabra pra editar.',
        })

    try:
        data = json.loads(request.body)
    except Exception:
        return JsonResponse({'ok': False, 'mensagem': 'Payload inválido.'})

    campo = data.get('campo')
    valor = data.get('valor')

    if campo not in ('minutos', 'valor_base', 'bonus_cheio', 'bonus_percent', 'marcado'):
        return JsonResponse({'ok': False, 'mensagem': f'Campo inválido: {campo}'})

    try:
        if campo == 'marcado':
            item.marcado = bool(valor)
        elif campo == 'bonus_percent':
            if valor not in ('0%', '50%', '100%'):
                return JsonResponse({'ok': False, 'mensagem': 'Percentual inválido.'})
            item.bonus_percent = valor
        else:
            setattr(item, campo, Decimal(str(valor)))

        item.save()
        item.repasse.recalcular_totais()

    except Exception as e:
        return JsonResponse({'ok': False, 'mensagem': str(e)})

    return JsonResponse({
        'ok': True,
        'item': _serializar_item(item),
        'repasse': _serializar_repasse(item.repasse),
    })


# =========================================================
# MARCAR / DESMARCAR TODAS
# =========================================================
@login_required
@require_POST
def marcar_todas(request, pk):
    repasse = get_object_or_404(Repasse, pk=pk, usuario_web=request.user)

    if _esta_finalizado(repasse):
        return JsonResponse({
            'ok': False,
            'mensagem': 'Esse repasse está finalizado. Reabra pra editar.',
        })

    try:
        data = json.loads(request.body)
    except Exception:
        return JsonResponse({'ok': False, 'mensagem': 'Payload inválido.'})

    marcar = bool(data.get('marcar', True))
    repasse.itens.update(marcado=marcar)
    repasse.recalcular_totais()

    return JsonResponse({
        'ok': True,
        'itens': [_serializar_item(i) for i in repasse.itens.all()],
        'repasse': _serializar_repasse(repasse),
    })


# =========================================================
# ARREDONDAR
# =========================================================
@login_required
@require_POST
def arredondar(request, pk):
    repasse = get_object_or_404(Repasse, pk=pk, usuario_web=request.user)

    if _esta_finalizado(repasse):
        return JsonResponse({
            'ok': False,
            'mensagem': 'Esse repasse está finalizado. Reabra pra editar.',
        })

    try:
        data = json.loads(request.body)
    except Exception:
        return JsonResponse({'ok': False, 'mensagem': 'Payload inválido.'})

    modo = data.get('modo', 'exato')
    if modo not in ('baixo', 'cima', 'exato'):
        return JsonResponse({'ok': False, 'mensagem': 'Modo inválido.'})

    repasse_service.arredondar_repasse(repasse, modo)

    return JsonResponse({
        'ok': True,
        'itens': [_serializar_item(i) for i in repasse.itens.all()],
        'repasse': _serializar_repasse(repasse),
    })


# =========================================================
# RECALCULAR
# =========================================================
@login_required
@require_POST
def recalcular(request, pk):
    repasse = get_object_or_404(Repasse, pk=pk, usuario_web=request.user)

    if _esta_finalizado(repasse):
        return JsonResponse({
            'ok': False,
            'mensagem': 'Esse repasse está finalizado. Reabra pra editar.',
        })

    repasse_service.recalcular_repasse(repasse)
    messages.success(request, 'Repasse recalculado.')

    return JsonResponse({
        'ok': True,
        'itens': [_serializar_item(i) for i in repasse.itens.all()],
        'repasse': _serializar_repasse(repasse),
    })


# =========================================================
# ESCOLHER MINUTO (resolver ambiguidade)
# =========================================================
@login_required
@require_POST
def escolher_minuto(request, pk):
    from ..services.repasse_service import _normalizar

    item = get_object_or_404(
        ItemRepasse, pk=pk, repasse__usuario_web=request.user
    )

    if _esta_finalizado(item.repasse):
        return JsonResponse({
            'ok': False,
            'mensagem': 'Esse repasse está finalizado. Reabra pra editar.',
        })

    repasse = item.repasse

    try:
        data = json.loads(request.body)
    except Exception:
        return JsonResponse({'ok': False, 'mensagem': 'Payload inválido.'})

    minutos = data.get('minutos')
    if minutos is None:
        return JsonResponse({'ok': False, 'mensagem': 'Minutos não informados.'})

    try:
        minutos_dec = Decimal(str(minutos))
    except Exception:
        return JsonResponse({'ok': False, 'mensagem': 'Minuto inválido.'})

    esp_norm = _normalizar(item.especialidade)
    base = _extrair_base(esp_norm)

    candidatas = [
        it for it in repasse.itens.all()
        if _extrair_base(_normalizar(it.especialidade)) == base
    ]

    marcadas = [it for it in candidatas if it.marcado]
    alvo = marcadas if marcadas else candidatas

    if not alvo:
        return JsonResponse({
            'ok': False,
            'mensagem': 'Nenhuma linha encontrada para essa especialidade.',
        })

    afetados = 0
    for it in alvo:
        it.minutos = minutos_dec
        it.faltou_regra = False
        it.save()
        afetados += 1

    repasse.recalcular_totais()

    return JsonResponse({
        'ok': True,
        'afetados': afetados,
        'marcadas': len(marcadas),
        'total_candidatas': len(candidatas),
        'itens': [_serializar_item(i) for i in repasse.itens.all()],
        'repasse': _serializar_repasse(repasse),
    })


def _extrair_base(esp_norm):
    if ' - ' in esp_norm:
        return esp_norm.split(' - ')[0].strip()
    return esp_norm.strip()


# =========================================================
# ZERAR
# =========================================================
@login_required
@require_POST
def zerar_valores(request, pk):
    repasse = get_object_or_404(Repasse, pk=pk, usuario_web=request.user)

    if _esta_finalizado(repasse):
        return JsonResponse({
            'ok': False,
            'mensagem': 'Esse repasse está finalizado. Reabra pra editar.',
        })

    for item in repasse.itens.all():
        item.minutos = Decimal('0')
        item.valor_base = Decimal('0')
        item.bonus_cheio = Decimal('0')
        item.save()

    repasse.recalcular_totais()

    return JsonResponse({
        'ok': True,
        'itens': [_serializar_item(i) for i in repasse.itens.all()],
        'repasse': _serializar_repasse(repasse),
    })


# =========================================================
# FINALIZAR / REABRIR
# =========================================================
@login_required
@require_POST
def finalizar_repasse(request, pk):
    repasse = get_object_or_404(Repasse, pk=pk, usuario_web=request.user)

    if repasse.status == 'finalizado':
        return JsonResponse({
            'ok': False,
            'mensagem': 'Esse repasse já está finalizado.',
        })

    marcados = repasse.itens.filter(marcado=True).count()
    if marcados == 0:
        return JsonResponse({
            'ok': False,
            'mensagem': 'Marque pelo menos uma linha antes de finalizar.',
        })

    repasse.recalcular_totais()

    repasse.status = 'finalizado'
    repasse.finalizado_em = timezone.now()
    repasse.save(update_fields=['status', 'finalizado_em', 'atualizado_em'])

    return JsonResponse({
        'ok': True,
        'mensagem': f'Repasse finalizado com {marcados} itens.',
        'repasse': _serializar_repasse(repasse),
    })


@login_required
@require_POST
def reabrir_repasse(request, pk):
    repasse = get_object_or_404(Repasse, pk=pk, usuario_web=request.user)

    if repasse.status != 'finalizado':
        return JsonResponse({
            'ok': False,
            'mensagem': 'Esse repasse não está finalizado.',
        })

    repasse.status = 'rascunho'
    repasse.finalizado_em = None
    repasse.save(update_fields=['status', 'finalizado_em', 'atualizado_em'])

    return JsonResponse({
        'ok': True,
        'mensagem': 'Repasse reaberto para edição.',
        'repasse': _serializar_repasse(repasse),
    })


# =========================================================
# EXPORTAR EXCEL (do repasse individual)
# =========================================================
@login_required
def exportar_excel(request, pk):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    repasse = get_object_or_404(Repasse, pk=pk, usuario_web=request.user)
    marcados = repasse.itens.filter(marcado=True)

    if not marcados.exists():
        messages.warning(request, 'Nenhuma linha marcada.')
        return redirect('siresp_app:repasse_ver', pk=pk)

    wb = Workbook()
    ws = wb.active
    ws.title = "Repasse"

    ws["A1"] = "Repasse Médico"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:N1")
    ws["A1"].alignment = Alignment(horizontal="center")

    ext = repasse.extracao
    ws["A2"] = (
        f"Profissional: {ext.medico_nome}  |  "
        f"Período: {ext.data_ini} a {ext.data_fim}"
    )
    ws.merge_cells("A2:N2")
    ws["A2"].alignment = Alignment(horizontal="center")
    ws["A2"].font = Font(italic=True, size=11)

    linha_cab = 4
    cabecalhos = [
        "Especialidade", "Grupo",
        "Oferta", "Atendimentos", "Minutos",
        "Valor Base (R$)", "Bônus 100% (R$)", "Bônus %", "Bônus Aplicado (R$)",
        "Valor/Hora (R$)",
        "Horas Real", "Horas Final",
        "Total Real (R$)", "Total Final (R$)",
    ]
    fonte_cab = Font(bold=True, color="FFFFFF")
    fundo_cab = PatternFill(start_color="003366", end_color="003366", fill_type="solid")

    for col, texto in enumerate(cabecalhos, start=1):
        cell = ws.cell(row=linha_cab, column=col, value=texto)
        cell.font = fonte_cab
        cell.fill = fundo_cab
        cell.alignment = Alignment(horizontal="center", vertical="center")

    larguras = [45, 22, 10, 13, 10, 15, 15, 10, 16, 15, 12, 12, 15, 15]
    for i, w in enumerate(larguras):
        col_letter = chr(ord('A') + i)
        ws.column_dimensions[col_letter].width = w

    linha = linha_cab + 1
    for item in marcados:
        ws.cell(row=linha, column=1, value=item.especialidade)
        ws.cell(row=linha, column=2, value=item.grupo_nome)
        ws.cell(row=linha, column=3, value=item.oferta)
        ws.cell(row=linha, column=4, value=item.atendimentos)
        ws.cell(row=linha, column=5, value=float(item.minutos))
        ws.cell(row=linha, column=6, value=float(item.valor_base))
        ws.cell(row=linha, column=7, value=float(item.bonus_cheio))
        ws.cell(row=linha, column=8, value=item.bonus_percent)
        ws.cell(row=linha, column=9, value=float(item.bonus_aplicado))
        ws.cell(row=linha, column=10, value=float(item.valor_hora))
        ws.cell(row=linha, column=11, value=float(item.horas_real))
        ws.cell(row=linha, column=12, value=float(item.horas_final))
        ws.cell(row=linha, column=13, value=float(item.total_real))
        ws.cell(row=linha, column=14, value=float(item.total_final))

        ws.cell(row=linha, column=5).number_format = '0'
        for c in (6, 7, 9, 10, 13, 14):
            ws.cell(row=linha, column=c).number_format = 'R$ #,##0.00'
        for c in (11, 12):
            ws.cell(row=linha, column=c).number_format = '0.00'
        linha += 1

    ws.cell(row=linha, column=1, value="TOTAL GERAL").font = Font(bold=True)
    ws.cell(row=linha, column=11, value=float(repasse.horas_total_real)).font = Font(bold=True)
    ws.cell(row=linha, column=12, value=float(repasse.horas_total_final)).font = Font(bold=True)
    ws.cell(row=linha, column=13, value=float(repasse.valor_total_real)).font = Font(bold=True)
    ws.cell(row=linha, column=14, value=float(repasse.valor_total_final)).font = Font(bold=True)

    for c in (11, 12):
        ws.cell(row=linha, column=c).number_format = '0.00'
    for c in (13, 14):
        ws.cell(row=linha, column=c).number_format = 'R$ #,##0.00'

    thin = Side(border_style="thin", color="000000")
    for c in range(1, 15):
        cell = ws.cell(row=linha, column=c)
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    nome = f"repasse_{ext.medico_nome.replace(' ', '_')}_{ext.data_ini.replace('/', '-')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{nome}"'
    wb.save(response)
    return response


# =========================================================
# SERIALIZAÇÃO
# =========================================================
def _serializar_item(item):
    return {
        'id': item.id,
        'especialidade': item.especialidade,
        'oferta': item.oferta,
        'atendimentos': item.atendimentos,
        'minutos': float(item.minutos),
        'valor_base': float(item.valor_base),
        'bonus_cheio': float(item.bonus_cheio),
        'bonus_percent': item.bonus_percent,
        'bonus_aplicado': float(item.bonus_aplicado),
        'valor_hora': float(item.valor_hora),
        'horas_real': float(item.horas_real),
        'horas_final': float(item.horas_final),
        'total_real': float(item.total_real),
        'total_final': float(item.total_final),
        'grupo_nome': item.grupo_nome,
        'marcado': item.marcado,
        'faltou_regra': item.faltou_regra,
    }


def _serializar_repasse(repasse):
    return {
        'id': repasse.id,
        'horas_total_real': float(repasse.horas_total_real),
        'horas_total_final': float(repasse.horas_total_final),
        'valor_total_real': float(repasse.valor_total_real),
        'valor_total_final': float(repasse.valor_total_final),
        'status': repasse.status,
        'finalizado': repasse.status == 'finalizado',
    }